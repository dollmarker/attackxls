# -*- coding: utf-8 -*-
"""
xls.py / idharvest.py —— .xls 附件全字段提取（不脱敏）

只处理老式 .xls（OLE2 复合文档 / BIFF8）

用法:
    python xls.py <文件或目录> [更多路径...]
    python xls.py            # 不带参数时扫当前目录

对每个 .xls，在它旁边生成:
    <同名>_数据.txt          完整表格，TSV（Tab 分隔），原样保留所有列
    <同名>_数据.xlsx         同一份数据，一个来源表一个 sheet，便于阅读
    <同名>_身份证号.txt      仅 18 位身份证号（用于计数取证）
    <同名>_身份证号_15位.txt  老式 15 位号（有才生成）

原理: .xls 不是 zip，是 OLE2 复合文档（CFB），魔数 D0CF11E0A1B11AE1，**不能 unzip**。
  内部流：Workbook（旧版 Book）/ SummaryInformation / DocumentSummaryInformation
  （WPS 另写 ETExtData）。单元格数据在 BIFF8 记录序列里（[类型 2B][长度 2B][数据]），
  文本集中在 SST 记录（0x00FC），跨块时会被拆进 CONTINUE（0x003C）—— 自己解析极易翻车，
  所以单元格用 xlrd 取，再用裸字节正则兜底扫描（两法取并集）。

 关键：数据不一定在单元格里。
  实测坑：某些 .xls（业务系统导出后经 WPS 转存）把**整份名册**塞在
  **BIFF 记录 0x005A（MS-XLS 的 Crn 记录）**里，工作表里根本没有"身份证号"这一列。
  这类载荷 xlrd / openpyxl 一律看不见 —— 只按单元格导出就会**静默丢数据**，
  症状正是"号码.txt 里有 215 个，_数据.xlsx 里却找不到身份证列"。
  所以本脚本加了记录级扫描：解码 0x005A 载荷 → 推断表头 → 作为独立一节
  并入 _数据.txt / _数据.xlsx，并做"单元格 vs 非单元格"的载体对照。
"""
import os
import re
import struct
import sys

# ---------- 身份证号 ----------
# 前 2 位限定为合法行政区划码，顺带滤掉字节扫描产生的大部分噪声
P = r"(?:1[1-5]|2[1-3]|3[1-7]|4[1-6]|5[0-4]|6[1-5]|71|81|82)"
MD = r"(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])"          # 月 + 日（真实日期，用来挡准考证号）
RE18_T = re.compile(r"(?<!\d)(" + P + r"\d{15}[\dXx])(?!\d)")
RE18_B = re.compile((r"(?<!\d)(" + P + r"\d{15}[\dXx])(?!\d)").encode())
RE15_T = re.compile(r"(?<!\d)(" + P + r"\d{4}\d{2}" + MD + r"\d{3})(?!\d)")   # 6位地区+YY+MMDD+顺序
RE15_B = re.compile((r"(?<!\d)(" + P + r"\d{4}\d{2}" + MD + r"\d{3})(?!\d)").encode())
RE_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
RE_CJK = re.compile(r"[\u4e00-\u9fff]{2,8}")

XLS_EXT = {".xls"}
# 本脚本自己的产物，扫描目录时跳过，免得把上次的输出当输入再跑一遍
OWN_OUTPUT = re.compile(r"(_数据|_身份证号(_15位)?)\.(xlsx|txt)$", re.I)

WEIGHTS = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
CHECK = "10X98765432"

XLSX_MAX_ROWS = 200000        # 超过就不生成 xlsx，避免内存爆炸

CFB_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
CRN_REC = 0x005A              # MS-XLS: Crn —— 实测的"非单元格载体"


def checksum_ok(s):
    """GB 11643-1999 / ISO 7064 MOD 11-2 校验位 —— 证明数据是真的，不是随机噪声"""
    s = s.upper()
    if len(s) != 18 or not s[:17].isdigit():
        return False
    return CHECK[sum(int(c) * w for c, w in zip(s[:17], WEIGHTS)) % 11] == s[17]


def _clean(s):
    return (s or "").replace("\t", " ").replace("\r", " ").replace("\n", " ").strip()


# ---------- OLE2 / CFB（自带解析，不依赖 olefile） ----------
def cfb_stream(data, want="Workbook"):
    """极简 CFB 读取器：取出某个流的原始字节。失败返回 None。"""
    if len(data) < 512 or data[:8] != CFB_MAGIC:
        return None
    try:
        sec = 1 << struct.unpack_from("<H", data, 30)[0]
        if sec not in (512, 4096):
            return None
        n_fat = struct.unpack_from("<I", data, 44)[0]
        dir_start = struct.unpack_from("<I", data, 48)[0]
        mini_cut = struct.unpack_from("<I", data, 56)[0]
        mini_start = struct.unpack_from("<I", data, 60)[0]
        difat_start = struct.unpack_from("<I", data, 68)[0]
    except Exception:
        return None

    def sector(n):
        off = 512 + n * sec
        return data[off:off + sec]

    per = sec // 4
    fat = []
    # 头部 DIFAT 只放 109 项，每项指向一个 FAT 扇区
    for i in range(min(n_fat, 109)):
        s = struct.unpack_from("<I", data, 76 + i * 4)[0]
        if s >= 0xFFFFFFF0:
            break
        blk = sector(s)
        if len(blk) < sec:
            break
        fat += list(struct.unpack_from("<%dI" % per, blk, 0))
    nxt, guard = difat_start, 0
    while nxt < 0xFFFFFFF0 and guard < 65536:          # 超长 FAT 的 DIFAT 链
        blk = sector(nxt)
        if len(blk) < sec:
            break
        ents = list(struct.unpack_from("<%dI" % per, blk, 0))
        for e in ents[:-1]:
            if e >= 0xFFFFFFF0:
                continue
            b2 = sector(e)
            if len(b2) < sec:
                continue
            fat += list(struct.unpack_from("<%dI" % per, b2, 0))
        nxt = ents[-1]
        guard += 1

    def chain(start):
        out, cur, g = [], start, 0
        while cur < 0xFFFFFFF0 and g < 2000000:
            out.append(cur)
            cur = fat[cur] if cur < len(fat) else 0xFFFFFFFE
            g += 1
        return out

    def read(start, size):
        return b"".join(sector(n) for n in chain(start))[:size]

    dirdata = read(dir_start, 1 << 30)
    ents = []
    for i in range(len(dirdata) // 128):
        e = dirdata[i * 128:(i + 1) * 128]
        nlen = struct.unpack_from("<H", e, 64)[0]
        typ = e[66]
        if typ not in (1, 2, 5) or nlen < 2:
            continue
        ents.append((e[:nlen - 2].decode("utf-16-le", "ignore"), typ,
                     struct.unpack_from("<I", e, 116)[0],
                     struct.unpack_from("<Q", e, 120)[0]))

    target = next((x for x in ents if x[1] == 2 and x[0] == want), None)
    if target is None:
        return None
    _, _, start, size = target
    if size == 0:
        return b""
    root = next((x for x in ents if x[1] == 5), None)
    if size >= mini_cut or root is None:
        return read(start, size)

    # 小流（< mini_cut）放在 Root Entry 的 mini stream 里，另有一套 64B 粒度 FAT
    mini = read(root[2], root[3])
    mfat = []
    for s in chain(mini_start):
        blk = sector(s)
        mfat += list(struct.unpack_from("<%dI" % per, blk, 0))
    out, cur, g = b"", start, 0
    while cur < 0xFFFFFFF0 and g < 2000000:
        out += mini[cur * 64:(cur + 1) * 64]
        cur = mfat[cur] if cur < len(mfat) else 0xFFFFFFFE
        g += 1
    return out[:size]


# ---------- .xls 单元格解析 ----------
def xls_sections(data):
    """用 xlrd 逐单元格读，返回 [(表名, 二维表), ...]"""
    out = []
    try:
        import xlrd
        bk = xlrd.open_workbook(file_contents=data)
    except Exception:
        return out
    for sh in bk.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                ct = sh.cell_type(r, c)
                v = sh.cell_value(r, c)
                if ct == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate_as_datetime(v, bk.datemode).strftime("%Y-%m-%d")
                    except Exception:
                        pass
                elif ct == xlrd.XL_CELL_NUMBER and isinstance(v, float) and v.is_integer():
                    v = int(v)
                row.append("" if v is None else str(v))
            if any(x.strip() for x in row):
                rows.append(row)
        if rows:
            out.append((sh.name, rows))
    return out


def scan_bytes(data):
    """裸字节兜底：ASCII 单字节 + UTF-16LE 宽字符两路都扫"""
    a, b = set(), set()
    for m in RE18_B.finditer(data):
        a.add(m.group(1).decode("ascii").upper())
    for m in RE15_B.finditer(data):
        b.add(m.group(1).decode("ascii"))
    try:
        txt = data.decode("utf-16-le", "ignore")
        a |= {m.group(1).upper() for m in RE18_T.finditer(txt)}
        b |= {m.group(1) for m in RE15_T.finditer(txt)}
    except Exception:
        pass
    return a, b


def ids_in(rows):
    """从二维表里逐格精确提取号码"""
    i18, i15 = set(), set()
    for row in rows:
        for v in row:
            t = str(v)
            i18 |= {m.group(1).upper() for m in RE18_T.finditer(t)}
            i15 |= {m.group(1) for m in RE15_T.finditer(t)}
    return i18, i15


def ids_of_sections(sections):
    i18, i15 = set(), set()
    for _, rows in sections:
        a, b = ids_in(rows)
        i18 |= a
        i15 |= b
    return i18, i15


# ---------- 非单元格载体：BIFF 记录级扫描 ----------
def biff_records(blob):
    """按 [类型 2B][长度 2B][数据] 切分 BIFF 记录"""
    out, i, n = [], 0, len(blob)
    while i + 4 <= n:
        rt, rl = struct.unpack_from("<HH", blob, i)
        if rt == 0 or i + 4 + rl > n:
            break
        out.append((i, rt, rl))
        i += 4 + rl
    return out


CRN_MAX_SKIP = 24            # 单条记录内允许跳过的对齐填充字节上限


def _crn_strict(d, o, col, col_last):
    """从 o 起按 [tag 1B][len 2B][grbit 1B][数据] 连续解析到 col_last。

    -> (列号->值, 结束偏移, clean)
       clean=False 表示停在 o 处时字段头越界（说明中间夹了填充，需要重新对齐）。
    """
    cells = {}
    while col <= col_last:
        if o + 4 > len(d):
            return cells, o, False
        ln = struct.unpack_from("<H", d, o + 1)[0]
        flag = d[o + 3] & 1
        nb = ln * 2 if flag else ln        # grbit&1: 1=UTF-16LE, 0=单字节压缩
        if o + 4 + nb > len(d):
            return cells, o, False
        cells[col] = d[o + 4:o + 4 + nb].decode("utf-16-le" if flag else "latin1", "replace")
        col += 1
        o += 4 + nb
    return cells, o, True


def _crn_resync(d, o, col, col_last):
    """字段头越界时，在 o 之后的有界窗口里找下一个能"恰好吃完整条记录"的对齐点。

    实测有的行会在字段之间夹 5 字节对齐填充（如序号 66），没有这一步整行会丢。
    """
    hi = min(o + 1 + CRN_MAX_SKIP, len(d))
    best = None
    for k in range(o + 1, hi):
        cells, end, _ = _crn_strict(d, k, col, col_last)
        if not cells:
            continue
        if end == len(d):
            return k                       # 首选：正好收尾
        if best is None:
            best = k
    return best


def decode_crn(d):
    """解码一条 Crn(0x005A) 载荷。

    头部布局（实测，注意不是"版本+序号"）：
        [colLast 1B][colFirst 1B][rw 2B]
        · 普通行 colFirst=0 / colLast=7 → 8 列，列 0..7
        · 字段多的行会被 Excel **拆成多条 Crn 记录，共用同一个 rw**：
          实测第 100 行 = (colFirst=0,colLast=5) + (colFirst=7,colLast=7)
          所以必须按 rw 归并、按 colFirst/colLast 定位列，否则不是丢行就是串列。

    字段体 [tag 1B][len 2B][grbit 1B][数据]，grbit 是 BIFF8 压缩标志：
        grbit & 1 == 1 → UTF-16LE，占 len*2 字节（中文姓名/民族/住址）
        grbit & 1 == 0 → 单字节压缩存储，占 len 字节（纯数字的身份证/考生号/学号）
        ★ 老代码把 grbit==0 当成"空字符串"，把压缩存储的那几列整列丢掉 ——
          121 个号码一个都读不出来。只有 len==0 才是真·空字段。

    -> (rw, {列号: 值}) 或 None
    """
    if len(d) < 5:
        return None
    col_last, col_first, rw = d[0], d[1], struct.unpack_from("<H", d, 2)[0]
    if col_first > col_last:                       # 空行
        return (rw, {})
    row, o, col, skipped = {}, 4, col_first, 0
    while o < len(d):
        cells, end, clean = _crn_strict(d, o, col, col_last)
        if cells:
            row.update(cells)
            col += len(cells)
        o = end
        if o == len(d) or col > col_last:
            break
        if not clean and len(d) - o <= 3:          # 尾部 1~3 字节对齐填充
            break
        k = _crn_resync(d, o, col, col_last)
        if k is None:
            return (rw, row) if row else None
        skipped += k - o
        if skipped > CRN_MAX_SKIP:
            return (rw, row) if row else None
        o = k
    return (rw, row) if row else None


def guess_header(rows, ncols):
    """按列内容猜表头 —— 只给有把握的列命名，其余保留 字段N 占位（不丢列）"""
    head = []
    for c in range(ncols):
        vals = [r[c] for r in rows if c < len(r) and r[c]]
        uniq = set(vals)
        if not vals:
            head.append("字段%d" % (c + 1))
        elif all(RE18_T.fullmatch(v) or RE15_T.fullmatch(v) for v in uniq):
            head.append("身份证号")
        elif all(RE_DATE.fullmatch(v) for v in uniq):
            head.append("出生日期")
        elif uniq <= {"男", "女"}:
            head.append("性别")
        elif all(v.endswith("族") and len(v) <= 5 for v in uniq):
            head.append("民族")
        elif all(re.fullmatch(r"\d{14}", v) for v in uniq):
            head.append("考生号")
        elif all(re.fullmatch(r"\d{9}", v) for v in uniq):
            head.append("学号")
        elif any("身份证" in v for v in uniq):
            head.append("证件类型")
        elif uniq <= {"中国", "中华人民共和国"}:
            head.append("国籍")
        elif c == 0 and all(RE_CJK.fullmatch(v) for v in uniq):
            head.append("姓名")
        elif any(re.search(r"[省市区县镇乡村]", v) for v in uniq):
            head.append("家庭住址")
        elif any(re.search(r"\d+\s*-\s*\d+$", v) for v in uniq):
            head.append("专业班级")
        elif all(re.fullmatch(r"\d{6}", v) for v in uniq) and len(uniq) <= max(16, len(rows) // 3):
            head.append("专业代码")
        else:
            head.append("字段%d" % (c + 1))
    # 相邻规则：紧挨在"已命名列"左侧的纯数字短码列，多半就是它的代码列
    for c in range(ncols - 1):
        if head[c].startswith("字段") and not head[c + 1].startswith("字段"):
            vals = [r[c] for r in rows if c < len(r) and r[c]]
            if vals and all(re.fullmatch(r"\d{1,3}", v) for v in vals):
                head[c] = head[c + 1] + "代码"
    return head


def non_cell_sections(data, cell18):
    """扫非单元格载体。-> (额外分节, 诊断行, 藏在非单元格里的号码集合)"""
    blob = cfb_stream(data, "Workbook")
    if not blob:
        return [], [], set()
    recs = biff_records(blob)

    by_rw, bad = {}, 0
    nrec = 0
    for off, rt, rl in recs:
        if rt != CRN_REC:
            continue
        r = decode_crn(blob[off + 4:off + 4 + rl])
        if r is None:
            bad += 1
            continue
        rw, cells = r
        if cells:
            nrec += 1
            by_rw.setdefault(rw, {}).update(cells)   # 同一 rw 的多条记录按列号合并

    extra, lines = [], []
    hid18 = set()
    rows = []
    if by_rw:
        ncols = max(max(c) for c in by_rw.values()) + 1
        rows = [[c.get(i, "") for i in range(ncols)] for _, c in sorted(by_rw.items())]
        extra.append(("内嵌记录_Crn(0x%04X)" % CRN_REC, [guess_header(rows, ncols)] + rows))
        hid18, _ = ids_in(rows)

    missing = hid18 - cell18
    if missing:
        # 载体定位：这些号码到底落在哪种 BIFF 记录里（记录内文本多为 UTF-16LE，两路都要扫）
        stat = {}
        for off, rt, rl in recs:
            seg = blob[off + 4:off + 4 + rl]
            hit = {m.group(1).decode("ascii").upper() for m in RE18_B.finditer(seg)}
            try:
                hit |= {m.group(1).upper() for m in RE18_T.finditer(seg.decode("utf-16-le", "ignore"))}
            except Exception:
                pass
            cnt = len(hit & missing)
            if cnt:
                stat[rt] = stat.get(rt, 0) + cnt
        top = "、".join("0x%04X x%d" % (rt, c) for rt, c in sorted(stat.items(), key=lambda kv: -kv[1])[:3])
        lines.append("非单元格载体 0x%04X：%d 条记录 / 归并成 %d 行，其中 %d 个号码不在任何单元格里"
                     "（载体分布 %s）" % (CRN_REC, nrec, len(rows), len(missing), top or "?"))
        if nrec != len(rows):
            lines.append("  注：%d 条记录共用同一行号（字段过多被 Excel 拆行），已按行号归并" % (nrec - len(rows)))
        lines.append("  已把该载体解码成独立一节，并入 _数据.txt / _数据.xlsx（表头由内容推断）")
        if bad:
            lines.append("另有 %d 条 0x%04X 记录结构不符，已跳过" % (bad, CRN_REC))
    elif rows:
        lines.append("非单元格载体 0x%04X：%d 行（号码均已在单元格中出现）" % (CRN_REC, len(rows)))
    return extra, lines, hid18


def extract(data, name):
    """-> (sections, ids18, ids15, 诊断行)

    号码必须三路取并集，少一路就漏：
      · 裸字节扫描   —— 文本型号码（SST 里是 UTF-16/ASCII 字符）
      · 单元格提取   —— **数字型号码**在 BIFF 里是 IEEE double（RK/NUMBER 记录），
                        字节流里根本没有数字字符，只有 xlrd 把它转成字符串才看得见
      · 非单元格载体 —— 0x005A 之类记录里的号码
    """
    sec = xls_sections(data)
    cell18, cell15 = ids_of_sections(sec)
    extra, lines, hid18 = non_cell_sections(data, cell18)
    i18, i15 = scan_bytes(data)
    return sec + extra, i18 | cell18 | hid18, i15 | cell15, lines


# ---------- 输出 ----------
def write_tsv(path, sections):
    dst = os.path.splitext(path)[0] + "_数据.txt"
    # newline="" 才不会被二次翻译（用 newline="\r\n" 会把 \n 再换成 \r\n，每行多一个空行）
    with open(dst, "w", encoding="utf-8-sig", newline="") as f:
        for label, rows in sections:
            f.write("# ===== %s =====\r\n" % label)
            for row in rows:
                f.write("\t".join(_clean(str(v)) for v in row) + "\r\n")
            f.write("\r\n")
    return dst


def _writable(p):
    """目标文件是否可写（被 Excel/WPS 打开时会返回 False）"""
    if not os.path.exists(p):
        return True
    try:
        with open(p, "r+b"):
            return True
    except OSError:
        return False


def write_xlsx(path, sections):
    """-> (输出路径 或 None, 说明/失败原因)"""
    total = sum(len(r) for _, r in sections)
    if total > XLSX_MAX_ROWS:
        return None, "行数 %d 超过上限 %d" % (total, XLSX_MAX_ROWS)
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception as e:
        return None, "openpyxl 不可用（%s）" % e

    base = os.path.splitext(path)[0] + "_数据"
    dst, note = base + ".xlsx", ""
    if not _writable(dst):
        # 原文件正被 Excel/WPS 打开 —— 换个名字写出，绝不静默丢整节数据
        k = 2
        while not _writable("%s(%d).xlsx" % (base, k)):
            k += 1
            if k > 20:
                return None, "目标被占用且改名失败（请关闭 Excel/WPS 后重跑）"
        dst = "%s(%d).xlsx" % (base, k)
        note = "原文件被占用，已改名为 %s" % os.path.basename(dst)

    wb = Workbook(write_only=True)          # 流式写入：几万行也不会卡住
    used = set()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_align = Alignment(horizontal="center", vertical="center")
    for label, rows in sections:
        title = re.sub(r'[\[\]\:\*\?\/\\]', "_", label)[:31] or "Sheet"
        base_t, k = title, 1
        while title in used:
            k += 1
            title = "%s_%d" % (base_t[:28], k)
        used.add(title)
        ws = wb.create_sheet(title)
        for i, row in enumerate(rows):
            vals = [_clean(str(v)) for v in row]
            if i == 0:                       # 首行当表头
                cells = []
                for v in vals:
                    c = WriteOnlyCell(ws, value=v)
                    c.font, c.fill, c.alignment = head_font, head_fill, head_align
                    cells.append(c)
                ws.append(cells)
            else:
                ws.append(vals)
        if rows:
            ws.freeze_panes = "A2"
    try:
        wb.save(dst)
    except Exception as e:
        return None, "写入失败（%s）" % e
    return dst, note


def write_ids(path, ids, suffix=""):
    dst = os.path.splitext(path)[0] + "_身份证号" + suffix + ".txt"
    with open(dst, "w", encoding="utf-8", newline="\r\n") as f:
        for i in sorted(ids):
            f.write(i + "\n")
    return dst


def walk(target):
    if os.path.isfile(target):
        return [target]
    hits = []
    for root, _, files in os.walk(target):
        for fn in files:
            if OWN_OUTPUT.search(fn):
                continue
            if os.path.splitext(fn)[1].lower() in XLS_EXT:
                hits.append(os.path.join(root, fn))
    return sorted(hits)


def main(argv):
    files = []
    for t in (argv[1:] or ["."]):
        files += walk(t)

    g18, g15 = set(), set()
    for p in files:
        with open(p, "rb") as f:
            data = f.read()
        name = os.path.basename(p)
        if data[:8] != CFB_MAGIC:
            print("[-] %-40s 不是 OLE2/.xls，跳过" % name[:40])
            continue
        sec, i18, i15, diag = extract(data, name)
        g18 |= i18
        g15 |= i15

        nrows = sum(len(r) for _, r in sec)
        if not sec and not i18 and not i15:
            print("[-] %-40s 无表格、无身份证号" % name[:40])
            continue
        ok = sum(1 for i in i18 if checksum_ok(i))
        print("[+] %-40s 表 %d 张 / %d 行 | 18位 %5d (校验通过 %5d) | 15位 %3d"
              % (name[:40], len(sec), nrows, len(i18), ok, len(i15)))
        for ln in diag:
            print("      [!] %s" % ln)
        if sec:
            t = write_tsv(p, sec)
            print("      -> %s" % os.path.basename(t))
            x, why = write_xlsx(p, sec)
            if x:
                print("      -> %s%s" % (os.path.basename(x), ("（%s）" % why) if why else ""))
            else:
                print("      -> [x] 未生成 xlsx：%s" % why)
        if i18:
            print("      -> %s" % os.path.basename(write_ids(p, i18)))
        if i15:
            print("      -> %s" % os.path.basename(write_ids(p, i15, "_15位")))

    print("\n合计 %d 个文件；去重后 18 位 %d 个、15 位 %d 个"
          % (len(files), len(g18), len(g15)))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
